"""Create a short, nontechnical guide to the school Wi-Fi portal."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path(__file__).with_name("Simple_School_WiFi_Guide.docx")
GREEN = RGBColor(25, 77, 58)
MINT = RGBColor(52, 133, 90)
GRAY = RGBColor(97, 112, 103)
INK = RGBColor(38, 54, 44)


def shade(cell, fill):
    prop = cell._tc.get_or_add_tcPr()
    tag = OxmlElement("w:shd")
    tag.set(qn("w:fill"), fill)
    prop.append(tag)


def text(paragraph, value, *, bold=False, size=None, color=None):
    run = paragraph.add_run(value)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


def line(document, first, rest):
    paragraph = document.add_paragraph()
    text(paragraph, first, bold=True, color=GREEN)
    text(paragraph, rest)


def card(document, title, body, fill="EDF5EE"):
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    text(paragraph, title + "\n", bold=True, color=GREEN, size=11)
    text(paragraph, body, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def step_table(document):
    steps = [
        ("1", "Join the school Wi-Fi", "Connect your phone or laptop to the school's Wi-Fi network."),
        ("2", "Open the welcome page", "The school Wi-Fi page appears when you try to use the internet."),
        ("3", "Read and accept", "Read the Wi-Fi guidelines, tick the box, then select “Accept & connect.”"),
        ("4", "Start browsing", "When the school network confirms access, your website opens or you see “You're connected.”"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.55)
    table.columns[1].width = Inches(5.8)
    for index, (number, title, body) in enumerate(steps):
        row = table.add_row()
        for cell in row.cells:
            shade(cell, "F2F7F1" if index % 2 == 0 else "FAFBF8")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        text(row.cells[0].paragraphs[0], number, bold=True, color=MINT, size=18)
        paragraph = row.cells[1].paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(7)
        text(paragraph, title + "\n", bold=True, color=GREEN, size=11)
        text(paragraph, body, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.page_height = Inches(11.7)
section.page_width = Inches(8.3)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15
for name, size in (("Title", 29), ("Heading 1", 16), ("Heading 2", 12)):
    style = doc.styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = GREEN
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(7)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
text(header, "ILAYA NATIONAL HIGH SCHOOL  •  SCHOOL WI-FI", bold=True, color=MINT, size=8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
text(footer, "Simple guide  •  Page ", color=GRAY, size=8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_before = Pt(16)
text(paragraph, "QUICK GUIDE FOR STUDENTS, TEACHERS & STAFF", bold=True, color=MINT, size=9)
doc.add_heading("How Our School Wi-Fi\nPortal Works", 0)
paragraph = doc.add_paragraph()
text(paragraph, "The portal is the welcome page you see before using the school's internet. It asks you to read and accept the Wi-Fi guidelines.", size=11)
card(doc, "No username or password needed", "In the current setup, you only need to accept the guidelines. The school network then checks your device and allows it to connect.")

doc.add_heading("Connect in four easy steps", 1)
step_table(doc)

doc.add_heading("After you connect", 1)
line(doc, "Use the internet.  ", "You can visit websites for learning, teaching, and school work.")
line(doc, "Disconnect when finished.  ", "If you return to the portal, select “Disconnect this device” to end the session.")
card(doc, "Why do I see this page again?", "Your Wi-Fi session may have ended, your device may have reconnected, or the school network may need you to accept the guidelines again.", "F8F4E9")

doc.add_page_break()
doc.add_heading("What happens in the background?", 1)
doc.add_paragraph(
    "The school's Wi-Fi equipment connects your device to the network. The school gateway, called OPNsense, controls when the device may use the internet. "
    "Before you accept, internet access is limited. After you select “Accept & connect,” OPNsense confirms your access and lets your device browse."
)
card(doc, "A simple way to picture it", "Your device  →  School Wi-Fi  →  School gateway  →  Internet")

doc.add_heading("What the IT team can see", 1)
doc.add_paragraph(
    "The separate IT dashboard helps school staff check how busy the connection is, how many devices are connected, and whether monitored network equipment has problems. "
    "This helps them find connection issues faster."
)
card(doc, "About the current dashboard preview", "The dashboard currently shows sample information for demonstration. It will show the school's actual network activity only after IT connects it to OPNsense and Zabbix.", "F8F4E9")

doc.add_heading("If something does not work", 1)
for title, body in [
    ("The welcome page does not open: ", "Reconnect to the school Wi-Fi and try opening a website."),
    ("“Accept & connect” does not work: ", "Ask the school IT coordinator to check the Wi-Fi portal."),
    ("The internet is slow or unavailable: ", "Tell the IT coordinator where you are and when the problem started."),
]:
    line(doc, title, body)

doc.add_heading("Good to remember", 1)
doc.add_paragraph(
    "The portal page alone cannot turn on internet access. The school must set up its firewall and Wi-Fi equipment for the button to work. "
    "The Wi-Fi guidelines shown on the page are a draft until the school approves them."
)
paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_before = Pt(20)
text(paragraph, "For technical setup, see the project README and OPNsense template in the CaptivePortal GitHub repository.", color=GRAY, size=9)

doc.core_properties.title = "How Our School Wi-Fi Portal Works"
doc.core_properties.subject = "Simple guide for students, teachers, and school staff"
doc.core_properties.author = "Ilaya National High School project documentation"
doc.save(OUTPUT)
print(OUTPUT)
